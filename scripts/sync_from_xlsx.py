"""Refresh data/benchmarks.json from the Google Sheet, without hand-typing.

1. In the sheet: File > Download > Microsoft Excel (.xlsx)
2. python scripts/sync_from_xlsx.py path/to/Festive_Report_Raw_Data.xlsx
3. Run tests/e2e_test.py, commit, push. The host redeploys automatically.

Expects the same tab names and column order as 'Festive Report 2026 Raw Data'.
Categories are matched by name (emoji ignored); a new category name must be added to KEYS below.
"""
import json
import pathlib
import re
import sys

from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "benchmarks.json"
KEYS = {"fashionapparel": "fashion", "beautypersonalcare": "beauty", "homeliving": "home", "wellness": "wellness",
        "electronics": "electronics", "foodbeverages": "food", "toysgamessports": "toys", "babykids": "baby",
        "booksstationery": "books", "petsupplies": "pet", "hardwareauto": "hardware_auto", "others": "others"}
FIELDS = ["aov_inr", "cod_pct", "prepaid_pct", "ppcod_pct", "emi_pct", "cod_rto_pct", "prepaid_rto_pct",
          "blended_rto_pct", "cod_return_pct", "checkout_completion_pct", "cart_abandon_pct", "rebuy_days",
          "charm_price_inr", "charm_price_pct", "top10_gmv_share_pct", "tier3_orders_pct", "tier3_rto_pct",
          "festive_gmv_lift_pct", "festive_orders_lift_pct"]
INR = {"aov_inr", "charm_price_inr", "rebuy_days"}


def clean_name(s):
    s = re.sub(r"[^\w\s&,]", "", str(s)).strip()  # drop emoji
    return s, KEYS.get(re.sub(r"[^a-z]", "", s.lower()))


def num(v, field):
    """Handles both raw numbers (0.304 for 30.4% when the cell is %-formatted) and strings ('30.4%', '₹1,688', '−12%')."""
    if isinstance(v, (int, float)):
        return round(v * 100, 2) if field not in INR and abs(v) <= 1.5 else v
    s = str(v).replace("₹", "").replace(",", "").replace("%", "").replace("−", "-").replace("+", "").strip()
    return float(s)


def rows_after_header(ws, header_first_cell):
    seen = False
    for row in ws.iter_rows(values_only=True):
        if not seen:
            seen = row[0] == header_first_cell
            continue
        if row[0]:
            yield row


def main(xlsx):
    wb = load_workbook(xlsx, data_only=True)
    data = json.loads(OUT.read_text())  # keep meta + timeline as-is
    cats = {}
    for r in rows_after_header(wb["Category Benchmarks"], "Category"):
        name, key = clean_name(r[0])
        if not key:
            sys.exit(f"Unknown category '{name}': add it to KEYS in this script and to ALIASES in logic.py")
        c = {"key": key, "name": name}
        for i, f in enumerate(FIELDS):
            c[f] = num(r[i + 1], f)
        c["rebuy_days"] = int(c["rebuy_days"])
        c["rides_festive"] = str(r[20]).strip().lower().startswith("y")
        c["basket_risk"], c["top_products"] = [], []
        cats[key] = c
    for r in rows_after_header(wb["Basket Risk"], "Category"):
        _, key = clean_name(r[0])
        cats[key]["basket_risk"].append({"basket_size": r[1], "pct_of_orders": num(r[2], "x"),
                                         "cod_rto_pct": num(r[3], "x"), "riskiest": bool(r[4])})
    for r in rows_after_header(wb["What Sold"], "Category"):
        _, key = clean_name(r[0])
        cats[key]["top_products"].append({"rank": int(r[1]), "product": r[2], "pct_of_units": num(r[3], "x")})
    for c in cats.values():
        mix = c["cod_pct"] + c["prepaid_pct"] + c["ppcod_pct"]
        assert 99 <= mix <= 101, f"{c['name']}: payment mix sums to {mix}"
        assert sum(b["riskiest"] for b in c["basket_risk"]) == 1, f"{c['name']}: needs exactly one riskiest basket"
    data["categories"] = list(cats.values())
    OUT.write_text(json.dumps(data, indent=1, ensure_ascii=False))
    print(f"Updated {OUT} with {len(cats)} categories. Update meta.window in the JSON if the season changed.")


if __name__ == "__main__":
    main(sys.argv[1])
